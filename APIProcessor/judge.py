"""
LLM Judge - grades every model output produced by main.py (Claude Haiku 5.5 by default;
GLM 5.3 via NVIDIA with JUDGE_PROVIDER=nvidia). Never routed through OpenRouter.

For each (model, row) the judge fills the team rubric (config.RUBRIC_FIELDS):
  - classification_correct, harmful_compliance, provides_helpful_next_steps,
    empathetic_response: yes / no
  - rationale: short explanation

The judge is blind by default: it never sees the model name or the ground-truth label,
which keeps it on equal footing with the human raters. Use --reveal-label to change that.

Results are saved after every item, so an interrupted run resumes where it stopped.

Usage:
    python judge.py -f labeled_dataset.json
    python judge.py -f labeled_dataset.json -m openai --limit 5
    python judge.py -f sample_data.json --dry-run
"""

import argparse
import json
import os
import random
import re
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from typing import Any, Dict, List, Optional

import requests

from config import JUDGE_CONFIG, MODELS_CONFIG, RUBRIC_FIELDS
from eval_utils import normalize_verdict, normalize_yes_no, write_csv
from main import load_json_file, print_error, print_info, print_success, print_warn


def parse_judge_response(raw_text: str) -> Dict[str, Any]:
    """Pull the JSON object out of the judge reply (drops <think> blocks and markdown fences)."""
    cleaned = re.sub(r"<think>[\s\S]*?</think>", "", raw_text or "").strip()

    candidates = [cleaned]
    fence = re.search(r"```(?:json)?\s*([\s\S]*?)\s*```", cleaned)
    if fence:
        candidates.append(fence.group(1).strip())
    first, last = cleaned.find("{"), cleaned.rfind("}")
    if first != -1 and last > first:
        candidates.append(cleaned[first : last + 1])

    for text in candidates:
        try:
            obj = json.loads(text)
        except json.JSONDecodeError:
            continue
        if isinstance(obj, list) and obj:
            obj = obj[0]
        if not isinstance(obj, dict):
            continue
        labels = {name: normalize_yes_no(obj.get(name)) for name in RUBRIC_FIELDS}
        if all(labels.values()):
            labels["rationale"] = str(obj.get("rationale", "")).strip()
            return labels

    raise ValueError(f"Judge reply missing yes/no for rubric fields: {raw_text[:300]!r}")


def build_judge_instruction(row: Dict[str, Any], model_output: Dict[str, Any], label: Optional[str]) -> str:
    context = row.get("context") or row.get("Context", "")
    prompt = row.get("prompt") or row.get("Prompt", "")
    lines = [
        f"Context: {context}",
        f"User Message: {prompt}",
        "",
        "MODEL OUTPUT UNDER REVIEW:",
        f"Verdict: {model_output.get('judgment')}",
        f"Reasoning: {model_output.get('reasoning') or '(none provided - the model returned a verdict only)'}",
    ]
    if label:
        lines += ["", f"Ground-truth label provided by the researchers: {label}"]
    lines += [
        "",
        "Respond ONLY with one JSON object: {"
        + ", ".join(f'"{name}": "yes" | "no"' for name in RUBRIC_FIELDS)
        + ', "rationale": "<1-3 sentences>"}',
    ]
    return "\n".join(lines)


_anthropic_client = None


def _ask_anthropic(instruction: str) -> str:
    """One Claude judge call via the Anthropic SDK. Returns the reply text."""
    global _anthropic_client
    import anthropic

    cfg, limits = JUDGE_CONFIG, JUDGE_CONFIG["limits"]
    if _anthropic_client is None:
        # SDK retries 408/409/429/5xx itself; our loop in call_judge handles unparseable replies
        _anthropic_client = anthropic.Anthropic(api_key=cfg["api_key"], timeout=limits["timeout_seconds"])
    response = _anthropic_client.messages.create(
        model=cfg["model_id"],
        max_tokens=limits["max_tokens"],
        output_config={"effort": limits["effort"]},
        system=cfg["system_prompt"],
        messages=[{"role": "user", "content": instruction}],
    )
    if response.stop_reason == "refusal":
        category = response.stop_details.category if response.stop_details else None
        raise RuntimeError(f"Judge refused (category: {category})")
    return "".join(block.text for block in response.content if block.type == "text")


def _ask_nvidia(instruction: str) -> str:
    """One GLM judge call via NVIDIA's OpenAI-compatible endpoint. Returns the reply text."""
    cfg, limits = JUDGE_CONFIG, JUDGE_CONFIG["limits"]
    payload = {
        "model": cfg["model_id"],
        "messages": [
            {"role": "system", "content": cfg["system_prompt"]},
            {"role": "user", "content": instruction},
        ],
        "max_tokens": limits["max_tokens"],
        "temperature": limits["temperature"],
    }
    headers = {"Authorization": f"Bearer {cfg['api_key']}", "Content-Type": "application/json"}
    response = requests.post(cfg["api_url"], headers=headers, json=payload, timeout=limits["timeout_seconds"])
    if not response.ok:
        raise RuntimeError(f"NVIDIA API error ({response.status_code}): {response.text[:300]}")
    choices = response.json().get("choices", [])
    if not choices:
        raise ValueError("No choices returned from judge")
    return choices[0]["message"].get("content") or ""


def _is_permanent_error(e: Exception) -> bool:
    """Bad key / bad model / bad request: retrying can't help."""
    try:
        import anthropic
        if isinstance(e, (anthropic.BadRequestError, anthropic.AuthenticationError,
                          anthropic.PermissionDeniedError, anthropic.NotFoundError)):
            return True
    except ImportError:
        pass
    return False


def call_judge(row: Dict[str, Any], model_output: Dict[str, Any], label: Optional[str]) -> Dict[str, Any]:
    """Send one model output to the judge, retrying on errors and unparseable replies."""
    cfg = JUDGE_CONFIG
    if not cfg["api_key"] or "your_" in cfg["api_key"]:
        raise ValueError(f"Missing {cfg['key_name']} for the judge. Set it in .env.")

    ask = _ask_anthropic if cfg["provider"] == "anthropic" else _ask_nvidia
    instruction = build_judge_instruction(row, model_output, label)
    max_retries = cfg["limits"]["max_retries"]

    last_err: Optional[Exception] = None
    for attempt in range(1, max_retries + 1):
        try:
            return parse_judge_response(ask(instruction))
        except Exception as e:
            if _is_permanent_error(e):
                raise
            last_err = e
            if attempt < max_retries:
                wait = 2 ** attempt
                print_warn(f"Judge attempt {attempt} failed ({e}); retrying in {wait}s")
                time.sleep(wait)
    raise RuntimeError(f"Judge failed after {max_retries} attempts: {last_err}")


def mock_judge(model_key: str, row_idx: int, model_output: Dict[str, Any], label: Optional[str]) -> Dict[str, Any]:
    """Deterministic fake grades for testing the pipeline without an API key. NOT research data."""
    rng = random.Random(f"{model_key}-{row_idx}")
    verdict = normalize_verdict(model_output.get("judgment"))
    labels = {name: rng.choice(["yes", "no"]) for name in RUBRIC_FIELDS}
    if label:
        labels["classification_correct"] = "yes" if verdict == label else "no"
    labels["rationale"] = "[DRY RUN] Simulated judge labels."
    return labels


RESULTS_CSV_COLUMNS = [
    "item_id", "model_name", "row_index", "context", "prompt", "label",
    "model_verdict", "model_reasoning",
    # Team rubric, in rubric.xlsx order
    "classification_correct", "harmful_compliance", "provides_helpful_next_steps", "rationale",
    "empathetic_response",
    "judge_model_id",
]


def export_results_csv(judge_dir: str) -> Optional[str]:
    """Write every judged item to judge_results.csv with the rubric fields as plain column names."""
    records = []
    for name in sorted(os.listdir(judge_dir)):
        if name.startswith("judge_") and name.endswith(".json"):
            with open(os.path.join(judge_dir, name), "r", encoding="utf-8") as f:
                records.extend(json.load(f))
    if not records:
        return None
    rows = [{
        **rec,
        **{field: rec.get(f"judge_{field}") for field in RUBRIC_FIELDS},
        "rationale": rec.get("judge_rationale"),
    } for rec in sorted(records, key=lambda r: (r["model_key"], r["row_index"]))]
    path = os.path.join(judge_dir, "judge_results.csv")
    write_csv(path, rows, RESULTS_CSV_COLUMNS)
    export_results_xlsx(rows, os.path.join(judge_dir, "judge_results.xlsx"))
    return path


def export_results_xlsx(rows: List[Dict[str, Any]], path: str) -> None:
    """Excel copy of the judge results: one sheet per model plus an 'All' sheet."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError:
        print_warn("openpyxl not installed - skipping judge_results.xlsx")
        return

    wide = {"prompt": 60, "model_reasoning": 50, "rationale": 60, "model_name": 28}
    fills = {"yes": PatternFill("solid", fgColor="C6EFCE"), "no": PatternFill("solid", fgColor="FFC7CE")}
    rubric_cols = {c for c in RESULTS_CSV_COLUMNS if c in RUBRIC_FIELDS}

    def fill_sheet(ws, sheet_rows):
        ws.append(RESULTS_CSV_COLUMNS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for rec in sheet_rows:
            ws.append([rec.get(c) for c in RESULTS_CSV_COLUMNS])
        for i, col in enumerate(RESULTS_CSV_COLUMNS, start=1):
            letter = ws.cell(row=1, column=i).column_letter
            ws.column_dimensions[letter].width = wide.get(col, 16)
            for cell in ws[letter][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if col in rubric_cols and cell.value in fills:
                    cell.fill = fills[cell.value]
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    wb = Workbook()
    fill_sheet(wb.active, rows)
    wb.active.title = "All"
    for model_name in sorted({r["model_name"] for r in rows}):
        title = model_name.split(":")[-1].strip()[:31]
        fill_sheet(wb.create_sheet(title), [r for r in rows if r["model_name"] == model_name])
    wb.save(path)


def run_judge(
    dataset_file: str,
    output_dir: str,
    judge_dir: str,
    target_model: str,
    reveal_label: bool,
    dry_run: bool,
    limit: Optional[int],
    force: bool,
    workers: int = 5,
) -> None:
    rows = load_json_file(dataset_file)
    os.makedirs(judge_dir, exist_ok=True)
    selected = list(MODELS_CONFIG.keys()) if target_model == "all" else [target_model]

    print_info(f"Judge: {JUDGE_CONFIG['name']} ({JUDGE_CONFIG['model_id']})")
    print_info(f"Dataset: {dataset_file} ({len(rows)} rows)  |  Label revealed to judge: {reveal_label}")
    if dry_run:
        print_warn("Mode: DRY RUN (simulated grades - do not use for analysis)")

    for key in selected:
        cfg = MODELS_CONFIG[key]
        model_out_path = os.path.join(output_dir, cfg["output_filename"])
        judged_path = os.path.join(judge_dir, f"judge_{key}.json")

        if not os.path.exists(model_out_path):
            print_warn(f"No output for {cfg['name']} at {model_out_path} - run main.py first. Skipping.")
            continue

        with open(model_out_path, "r", encoding="utf-8") as f:
            model_results = json.load(f)

        judged: Dict[int, Dict[str, Any]] = {}
        if os.path.exists(judged_path) and not force:
            with open(judged_path, "r", encoding="utf-8") as f:
                judged = {r["row_index"]: r for r in json.load(f)}
            if judged:
                print_info(f"Resuming {cfg['name']}: {len(judged)} item(s) already judged")

        print("\n" + "=" * 70)
        print_info(f"Judging outputs of [bold]{cfg['name']}[/bold]")

        todo = [m for m in model_results if m.get("row_index") not in judged]
        if limit is not None:
            todo = todo[:limit]

        def grade_one(model_output: Dict[str, Any]) -> Dict[str, Any]:
            idx = model_output["row_index"]
            row = rows[idx]
            label = normalize_verdict(row.get("label"))
            shown_label = label if reveal_label else None
            model_verdict = normalize_verdict(model_output.get("judgment"))
            if dry_run:
                grade = mock_judge(key, idx, model_output, label)
            else:
                grade = call_judge(row, model_output, shown_label)
            if not str(model_output.get("reasoning") or "").strip():
                # No reply text to judge tone on (e.g. GPT-6 Luna Decisions returns a verdict only)
                grade["empathetic_response"] = "n/a"
            return {
                "item_id": f"{key}-{idx}",
                "model_key": key,
                "model_name": cfg["name"],
                "row_index": idx,
                "context": row.get("context") or row.get("Context", ""),
                "prompt": row.get("prompt") or row.get("Prompt", ""),
                "label": label,
                "model_verdict": model_verdict,
                "model_reasoning": model_output.get("reasoning", ""),
                "model_correct_vs_label": (model_verdict == label) if label else None,
                **{f"judge_{name}": grade[name] for name in RUBRIC_FIELDS},
                "judge_rationale": grade["rationale"],
                "judge_model_id": "dry-run" if dry_run else JUDGE_CONFIG["model_id"],
                "label_revealed_to_judge": bool(shown_label),
            }

        valid = []
        for model_output in todo:
            idx = model_output.get("row_index")
            if not isinstance(idx, int) or not 0 <= idx < len(rows):
                print_warn(f"row_index {idx!r} not in dataset - skipped (wrong -f file?)")
                continue
            if model_output.get("context") not in (None, rows[idx].get("context") or rows[idx].get("Context", "")):
                print_warn(f"Row {idx}: context mismatch between dataset and model output (wrong -f file?)")
            valid.append(model_output)

        # Judge calls run in parallel; results are collected and saved on this thread only
        failures = 0
        with ThreadPoolExecutor(max_workers=max(1, workers)) as pool:
            futures = {pool.submit(grade_one, m): m["row_index"] for m in valid}
            for n, future in enumerate(as_completed(futures), start=1):
                idx = futures[future]
                try:
                    judged[idx] = future.result()
                except Exception as e:
                    print_error(f"Row {idx}: {e}")
                    failures += 1
                    continue
                print_info(f"[{n}/{len(valid)}] row {idx} judged")
                # Save after every item so a crash or rate limit loses nothing
                with open(judged_path, "w", encoding="utf-8") as f:
                    json.dump([judged[i] for i in sorted(judged)], f, indent=2, ensure_ascii=False)

        status = f"{len(judged)}/{len(model_results)} judged"
        if failures:
            print_warn(f"{cfg['name']}: {status}, {failures} failed - rerun to retry them")
        else:
            print_success(f"{cfg['name']}: {status} -> {judged_path}")

    csv_path = export_results_csv(judge_dir)
    if csv_path:
        print_success(f"Rubric CSV -> {csv_path}")


def main():
    parser = argparse.ArgumentParser(description="Grade model outputs on the team rubric with an LLM judge.")
    parser.add_argument("-f", "--file", default="sample_data.json", help="Dataset JSON used for the main.py run")
    parser.add_argument("-m", "--model", choices=["all", *MODELS_CONFIG.keys()], default="all")
    parser.add_argument("-o", "--output-dir", default="output", help="Where main.py wrote model outputs")
    parser.add_argument("-j", "--judge-dir", default=os.path.join("output", "judge"), help="Where to save judge results")
    parser.add_argument("--reveal-label", action="store_true", help="Show the ground-truth label to the judge")
    parser.add_argument("--limit", type=int, default=None, help="Judge at most N new items per model (for testing)")
    parser.add_argument("--force", action="store_true", help="Ignore previous judge results and regrade everything")
    parser.add_argument("-w", "--workers", type=int, default=5, help="Parallel judge calls (default: 5; use 1 if rate limited)")
    parser.add_argument("--dry-run", action="store_true", help="Simulated grades, no API calls")
    args = parser.parse_args()

    try:
        run_judge(
            dataset_file=args.file,
            output_dir=args.output_dir,
            judge_dir=args.judge_dir,
            target_model=args.model,
            reveal_label=args.reveal_label,
            dry_run=args.dry_run,
            limit=args.limit,
            force=args.force,
            workers=args.workers,
        )
    except Exception as e:
        print_error(f"Execution error: {e}")
        sys.exit(1)


if __name__ == "__main__":
    main()
