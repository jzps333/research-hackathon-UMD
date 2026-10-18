"""
Shared helpers for the judge / human-comparison pipeline:
verdict normalization, CSV/XLSX table IO, and agreement statistics.
"""

import csv
import os
from typing import Any, Dict, List, Optional, Sequence

VERDICTS = ("Scam", "Not a scam")

_SCAM_WORDS = {"scam", "yes", "y", "true", "1", "fake", "fraud", "phishing"}
_NOT_SCAM_WORDS = {
    "not a scam", "not scam", "no", "n", "false", "0",
    "real", "legit", "legitimate", "benign", "safe",
}


def normalize_verdict(value: Any) -> Optional[str]:
    """Map loose labels (Excel cells, model output) to 'Scam' / 'Not a scam'. None if unknown/blank."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "Scam" if value else "Not a scam"
    v = str(value).strip().lower()
    if v in _SCAM_WORDS:
        return "Scam"
    if v in _NOT_SCAM_WORDS:
        return "Not a scam"
    return None


_YES_WORDS = {"yes", "y", "true", "1"}
_NO_WORDS = {"no", "n", "false", "0"}


def normalize_yes_no(value: Any) -> Optional[str]:
    """Map rubric cells ('Yes', 'y', True, 1...) to 'yes' / 'no'. None if blank or unknown."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "yes" if value else "no"
    v = str(value).strip().lower()
    if v in _YES_WORDS:
        return "yes"
    if v in _NO_WORDS:
        return "no"
    return None


# ------------------------------------------------------------------------------
# Table IO (CSV or Excel)
# ------------------------------------------------------------------------------

def read_table(path: str, sheet: Optional[str] = None) -> List[Dict[str, Any]]:
    """Read a .csv or .xlsx file into a list of dicts keyed by lower-cased, stripped header."""
    if not os.path.exists(path):
        raise FileNotFoundError(f"File not found: {path}")

    ext = os.path.splitext(path)[1].lower()
    if ext in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise ImportError("Reading Excel files requires openpyxl: pip install openpyxl")
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb.active
        rows = list(ws.iter_rows(values_only=True))
        wb.close()
        if not rows:
            return []
        headers = [str(h).strip().lower() if h is not None else "" for h in rows[0]]
        out = []
        for r in rows[1:]:
            if r is None or all(c is None or str(c).strip() == "" for c in r):
                continue
            out.append({headers[i]: r[i] for i in range(min(len(headers), len(r))) if headers[i]})
        return out

    # utf-8-sig strips the BOM Excel adds when saving CSV
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f)
        return [
            {(k or "").strip().lower(): v for k, v in row.items() if k}
            for row in reader
            if any((v or "").strip() for v in row.values() if isinstance(v, str))
        ]


FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def excel_safe(value: Any) -> Any:
    """Neutralize spreadsheet formula injection: prompts and model replies are untrusted text."""
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return "'" + value
    return value


def write_csv(path: str, rows: List[Dict[str, Any]], columns: Sequence[str]) -> None:
    """Write rows as CSV with a BOM so Excel opens UTF-8 text correctly."""
    os.makedirs(os.path.dirname(path) or ".", exist_ok=True)
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(columns), extrasaction="ignore")
        writer.writeheader()
        writer.writerows({k: excel_safe(v) for k, v in row.items()} for row in rows)


# ------------------------------------------------------------------------------
# Agreement statistics (pure python, no scipy)
# ------------------------------------------------------------------------------

def percent_agreement(a: Sequence[Any], b: Sequence[Any]) -> Optional[float]:
    if not a:
        return None
    return 100.0 * sum(x == y for x, y in zip(a, b)) / len(a)


def cohen_kappa(a: Sequence[Any], b: Sequence[Any]) -> Optional[float]:
    """Cohen's kappa for two raters on nominal labels. None when undefined (no variation)."""
    n = len(a)
    if n == 0:
        return None
    cats = set(a) | set(b)
    po = sum(x == y for x, y in zip(a, b)) / n
    pe = sum((list(a).count(c) / n) * (list(b).count(c) / n) for c in cats)
    if pe >= 1.0:
        return None
    return (po - pe) / (1 - pe)


def yes_rate(values: Sequence[str]) -> Optional[float]:
    """% of 'yes' among yes/no values; anything else (e.g. 'n/a') is left out."""
    values = [v for v in values if v in ("yes", "no")]
    if not values:
        return None
    return 100.0 * sum(v == "yes" for v in values) / len(values)
